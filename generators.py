# -*- coding: utf-8 -*-
import os
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import config

# ===== Цветовые константы =====
# DOCX цвета (RGB)
COLOR_SUBTITLE = RGBColor(0x71, 0x80, 0x96)   # Серый для подзаголовков
COLOR_ADDRESS = RGBColor(0x4A, 0x55, 0x68)    # Тёмно-серый для адреса
COLOR_ACCENT = RGBColor(0xD2, 0x26, 0x30)     # Красный акцент
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # Белый
COLOR_TABLE_HEADER_BG = "#F7FAFC"              # Светло-серый фон таблицы
COLOR_TABLE_ROW_BG = "#F8FAFC"                 # Фон строк таблицы
COLOR_SEPARATOR = RGBColor(0xE2, 0xE8, 0xF0)   # Цвет разделителя

# PDF цвета (ReportLab)
PDF_COLOR_GRAY = colors.HexColor('#718096')    # Серый для PDF
PDF_COLOR_DESC = colors.HexColor('#4A5568')    # Описание
PDF_COLOR_TEXT = colors.HexColor('#2D3748')    # Основной текст
PDF_COLOR_GRID = colors.HexColor('#CBD5E0')    # Сетка таблицы
PDF_COLOR_LIGHT_BG = colors.HexColor('#F7FAFC')# Светлый фон
PDF_COLOR_ACCENT = colors.HexColor('#D22630')  # Красный акцент для PDF

def set_cell_background(cell, hex_color):
    """Устанавливает цвет фона ячейки таблицы в Word."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def find_system_font(font_name, bold=False):
    """Ищет файл шрифта в системной папке Windows по имени. Возвращает полный путь или None."""
    fonts_dir = r"C:\Windows\Fonts"
    if not os.path.exists(fonts_dir):
        return None
    name_lower = font_name.lower().replace(' ', '')
    candidates = []
    for file in os.listdir(fonts_dir):
        if file.lower().endswith('.ttf'):
            file_stem = file.lower()[:-4].replace(' ', '')
            if bold:
                if name_lower in file_stem and ('bd' in file_stem or 'bold' in file_stem):
                    candidates.append(file)
            else:
                if name_lower in file_stem and 'bd' not in file_stem and 'bold' not in file_stem:
                    candidates.append(file)
    if candidates:
        return os.path.join(fonts_dir, candidates[0])
    for file in os.listdir(fonts_dir):
        if file.lower().endswith('.ttf'):
            file_stem = file.lower()[:-4].replace(' ', '')
            if name_lower in file_stem:
                if bold and ('bd' in file_stem or 'bold' in file_stem):
                    return os.path.join(fonts_dir, file)
                elif not bold and 'bd' not in file_stem and 'bold' not in file_stem:
                    return os.path.join(fonts_dir, file)
    return None

def get_cyrillic_font_registered(font_family=None):
    """Регистрирует шрифт в ReportLab и возвращает имена обычного и жирного шрифтов."""
    if font_family:
        reg_path = find_system_font(font_family, bold=False)
        bold_path = find_system_font(font_family, bold=True)
        if reg_path:
            try:
                pdfmetrics.registerFont(TTFont('UserFont', reg_path))
                if bold_path:
                    pdfmetrics.registerFont(TTFont('UserFontBold', bold_path))
                else:
                    pdfmetrics.registerFont(TTFont('UserFontBold', reg_path))
                return 'UserFont', 'UserFontBold'
            except Exception:
                pass
    arial_path = find_system_font("Arial", bold=False)
    arial_bold_path = find_system_font("Arial", bold=True)
    if arial_path:
        pdfmetrics.registerFont(TTFont('CyrillicFont', arial_path))
        pdfmetrics.registerFont(TTFont('CyrillicFont-Bold', arial_bold_path or arial_path))
        return 'CyrillicFont', 'CyrillicFont-Bold'
    return 'Helvetica', 'Helvetica-Bold'

def format_model_description(model):
    if model.get("description"):
        return model["description"]
    parts = []
    if model.get("manufacturer"):
        parts.append(f"Производитель: {model['manufacturer']}")
    if model.get("series"):
        parts.append(f"Серия/тип: {model['series']}")
    return "; ".join(parts) if parts else "Сплит-система настенного типа"

def hex_to_rgb(hex_color):
    """Преобразует HEX-строку в кортеж (R,G,B)."""
    hex_color = hex_color.lstrip('#')
    try:
        return int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    except:
        return 0x1B, 0x36, 0x5D

def generate_docx(variants, customer_data, filename):
    """Генерация коммерческого предложения в формате DOCX с настройками внешнего вида."""
    settings = config.load_settings()
    logo_path = settings.get("logo_path", "").strip()
    primary_hex = settings.get("primary_color", "#1B365D")
    font_name = settings.get("font_family", "Segoe UI")
    logo_width_cm = float(settings.get("logo_width", 3.0))
    
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = font_name
    for section in doc.sections:
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
    
    # --- ШАПКА ---
    header_table = doc.add_table(rows=2, cols=2)
    header_table.autofit = False
    header_table.columns[0].width = Cm(11.5)
    header_table.columns[1].width = Cm(5.7)
    
    left_cell = header_table.cell(0, 0)
    left_cell.text = settings.get("activity", "КОНДИЦИОНЕРЫ\nВЕНТИЛЯЦИЯ\nОТОПЛЕНИЕ")
    for para in left_cell.paragraphs:
        para.runs[0].bold = True
        para.runs[0].font.size = Pt(13)
        para.runs[0].font.color.rgb = COLOR_SUBTITLE
        para.runs[0].font.name = font_name
    
    right_top = header_table.cell(0, 1)
    if logo_path and os.path.exists(logo_path):
        right_top.text = ""
        p = right_top.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_img = p.add_run()
        run_img.add_picture(logo_path, width=Cm(logo_width_cm))
    else:
        right_top.text = ""
        p_logo = right_top.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_neva = p_logo.add_run("НЕВА")
        run_neva.bold = True
        run_neva.font.size = Pt(18)
        run_neva.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
        run_neva.font.name = font_name
        run_con = p_logo.add_run("КОН")
        run_con.bold = True
        run_con.font.size = Pt(18)
        run_con.font.color.rgb = COLOR_ACCENT
        run_con.font.name = font_name
    
    right_bottom = header_table.cell(1, 1)
    right_bottom.text = settings.get("address", "г. СПб, Детский пер., д. 5")
    right_bottom.paragraphs[0].runs[0].font.size = Pt(9.5)
    right_bottom.paragraphs[0].runs[0].font.color.rgb = COLOR_ADDRESS
    right_bottom.paragraphs[0].runs[0].font.name = font_name
    right_bottom.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(5)
    p_line.paragraph_format.space_after = Pt(10)
    run_line = p_line.add_run("—" * 90)
    run_line.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
    run_line.font.name = font_name
    
    p_contacts = doc.add_paragraph()
    run_c = p_contacts.add_run(settings.get("contacts", "тел. (812) 622-09-49   e-mail: mail@nevacon.ru"))
    run_c.font.size = Pt(9.5)
    run_c.font.color.rgb = COLOR_ADDRESS
    run_c.font.name = font_name
    
    p_web = doc.add_paragraph()
    p_web.paragraph_format.space_after = Pt(2)
    run_w = p_web.add_run(settings.get("website", "https://www.nevacon.ru"))
    run_w.font.size = Pt(9.5)
    run_w.bold = True
    run_w.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
    run_w.font.name = font_name
    
    p_serv = doc.add_paragraph()
    run_s = p_serv.add_run(settings.get("services", "проектирование – монтаж - сервисное обслуживание"))
    run_s.font.size = Pt(9.5)
    run_s.font.italic = True
    run_s.font.color.rgb = COLOR_SUBTITLE
    run_s.font.name = font_name
    
    doc.add_paragraph()
    
    # --- РЕКВИЗИТЫ ---
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    info_table.columns[0].width = Cm(3.0)
    info_table.columns[1].width = Cm(14.0)
    
    date_str = customer_data.get("date", datetime.now().strftime("%d.%m.%Y"))
    if re.match(r'\d{4}-\d{2}-\d{2}', date_str):
        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        date_str = date_obj.strftime("%d.%m.%Y")
    
    rows_data = [
        ("Дата:", date_str),
        ("Тел./факс:", customer_data.get("phone", "")),
        ("Кому:", customer_data.get("to", "")),
        ("От:", customer_data.get("from", "")),
        ("Объект:", customer_data.get("object", ""))
    ]
    
    for i, (label, value) in enumerate(rows_data):
        cell_lbl = info_table.cell(i, 0)
        cell_lbl.text = label
        cell_lbl.paragraphs[0].runs[0].bold = True
        cell_lbl.paragraphs[0].runs[0].font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
        cell_lbl.paragraphs[0].runs[0].font.name = font_name
        set_cell_background(cell_lbl, COLOR_TABLE_HEADER_BG)
        
        cell_val = info_table.cell(i, 1)
        cell_val.text = value
        cell_val.paragraphs[0].runs[0].font.name = font_name
    
    doc.add_paragraph()
    
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ")
    run.bold = True
    run.underline = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
    run.font.name = font_name
    
    doc.add_paragraph()
    
    intro_para = doc.add_paragraph(
        "Благодарим Вас за обращение в нашу фирму по вопросам вентиляции и "
        "кондиционирования воздуха. Для создания комфортных температурных условий "
        "в Вашем помещении предлагаем установить следующий вид кондиционеров:"
    )
    intro_para.runs[0].font.name = font_name
    doc.add_paragraph()
    
    # Загружаем модели ОДИН раз перед циклом (оптимизация)
    models = config.load_models()
    
    for idx, var in enumerate(variants, 1):
        m_key = var["model_key"]
        model = models.get(m_key)
        if not model:
            raise Exception(f"Модель '{m_key}' не найдена в базе данных.")
        
        h2 = doc.add_paragraph(f"{idx} вариант")
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(2)
        run_h2 = h2.add_run()
        run_h2.text = f"{idx} вариант"
        run_h2.font.size = Pt(13)
        run_h2.bold = True
        run_h2.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
        run_h2.font.name = font_name
        
        desc_text = format_model_description(model)
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(6)
        run_desc = p_desc.add_run(desc_text)
        run_desc.font.size = Pt(10.5)
        run_desc.font.italic = True
        run_desc.font.color.rgb = COLOR_ADDRESS
        run_desc.font.name = font_name
        
        table = doc.add_table(rows=2, cols=9)
        table.style = 'Table Grid'
        col_widths = [Cm(2.5), Cm(1.5), Cm(1.8), Cm(1.8), Cm(1.8), Cm(2.2), Cm(1.2), Cm(1.8), Cm(2.2)]
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w
        
        headers = [
            "Модель", "Уровень шума, дБ",
            "Произв. по холоду, кВт", "Произв. по теплу, кВт",
            "Потреб. мощн., кВт", "Габариты внутр. блока, мм",
            "Кол-во, шт.", "Цена, руб.", "Итого, руб."
        ]
        
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = h
            run_h = cell.paragraphs[0].runs[0]
            run_h.bold = True
            run_h.font.size = Pt(8.5)
            run_h.font.color.rgb = COLOR_WHITE
            run_h.font.name = font_name
            set_cell_background(cell, primary_hex.replace('#', ''))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        qty = var["quantity"]
        price = model["price"]
        total_equip = price * qty
        
        row = table.rows[1]
        row.cells[0].text = m_key
        row.cells[1].text = model["noise_level"]
        row.cells[2].text = model["cooling_kw"]
        row.cells[3].text = model["heating_kw"]
        row.cells[4].text = model["power_kw"]
        row.cells[5].text = model["dimensions"]
        row.cells[6].text = str(qty)
        row.cells[7].text = f"{price:,}".replace(",", " ")
        row.cells[8].text = f"{total_equip:,}".replace(",", " ")
        
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            cell.paragraphs[0].runs[0].font.name = font_name
            set_cell_background(cell, COLOR_TABLE_ROW_BG)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        doc.add_paragraph()
        
        tab_pos = Cm(17.27)
        
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
        run_inst = p.add_run("Монтаж кондиционера")
        run_inst.bold = True
        run_inst.font.name = font_name
        run_inst_val = p.add_run(f"\t{var['installation_price']:,} руб.".replace(",", " "))
        run_inst_val.font.name = font_name
        
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
        run_appr = p.add_run("Согласование места размещения наружного блока")
        run_appr.bold = True
        run_appr.font.name = font_name
        run_appr_val = p.add_run(f"\t{var['approval_price']:,} руб.".replace(",", " "))
        run_appr_val.font.name = font_name
        
        total = total_equip + var["installation_price"] + var["approval_price"]
        
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
        run_tot_lbl = p.add_run("ИТОГО")
        run_tot_lbl.bold = True
        run_tot_lbl.font.size = Pt(11)
        run_tot_lbl.font.color.rgb = COLOR_ACCENT
        run_tot_lbl.font.name = font_name
        run_tot_val = p.add_run(f"\t{total:,} руб.".replace(",", " "))
        run_tot_val.bold = True
        run_tot_val.font.size = Pt(11)
        run_tot_val.font.color.rgb = COLOR_ACCENT
        run_tot_val.font.name = font_name
        
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(8)
        p_sep.paragraph_format.space_after = Pt(8)
        run_sep = p_sep.add_run("—" * 90)
        run_sep.font.color.rgb = COLOR_SEPARATOR
        run_sep.font.name = font_name
    
    doc.add_paragraph()
    
    p_respect = doc.add_paragraph("С уважением,")
    p_respect.runs[0].font.name = font_name
    
    p_sig = doc.add_paragraph()
    run_sig = p_sig.add_run(settings.get("manager_signature", "Толстиков Игорь."))
    run_sig.bold = True
    run_sig.font.color.rgb = RGBColor(*hex_to_rgb(primary_hex))
    run_sig.font.name = font_name
    
    try:
        doc.save(filename)
    except PermissionError:
        raise PermissionError(f"Не удалось сохранить файл '{os.path.basename(filename)}'. Закройте его, если он открыт в Word, и попробуйте снова.")

def generate_pdf(variants, customer_data, filename):
    """Генерация коммерческого предложения в формате PDF с настройками внешнего вида."""
    settings = config.load_settings()
    logo_path = settings.get("logo_path", "").strip()
    primary_color_hex = settings.get("primary_color", "#1B365D")
    font_family = settings.get("font_family", "Segoe UI")
    logo_width_cm = float(settings.get("logo_width", 3.0))
    
    font_reg, font_bold_reg = get_cyrillic_font_registered(font_family)
    
    try:
        primary_color = colors.HexColor(primary_color_hex)
    except:
        primary_color = colors.HexColor("#1B365D")
    
    accent_color = colors.HexColor("#D22630")
    
    doc_tpl = SimpleDocTemplate(
        filename,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    style_normal = ParagraphStyle(
        'CyrNormal',
        parent=styles['Normal'],
        fontName=font_reg,
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#2D3748')
    )
    style_bold = ParagraphStyle(
        'CyrBold',
        parent=style_normal,
        fontName=font_bold_reg
    )
    style_right = ParagraphStyle(
        'CyrRight',
        parent=style_normal,
        alignment=2
    )
    style_title = ParagraphStyle(
        'CyrTitle',
        parent=style_normal,
        fontName=font_bold_reg,
        fontSize=16,
        leading=20,
        alignment=1,
        textColor=primary_color
    )
    style_h2 = ParagraphStyle(
        'CyrH2',
        parent=style_normal,
        fontName=font_bold_reg,
        fontSize=12,
        leading=15,
        spaceBefore=10,
        spaceAfter=2,
        textColor=primary_color
    )
    
    story = []
    
    left_html = settings.get("activity", "КОНДИЦИОНЕРЫ\nВЕНТИЛЯЦИЯ\nОТОПЛЕНИЕ").replace('\n', '<br/>')
    
    style_left_header = ParagraphStyle(
        'LeftHdr',
        parent=style_normal,
        fontName=font_bold_reg,
        fontSize=12,
        leading=14,
        textColor=colors.HexColor('#718096')
    )
    
    if logo_path and os.path.exists(logo_path):
        logo_img = Image(logo_path, width=logo_width_cm * 28.35)
        logo_table = Table([[logo_img]], colWidths=[265])
        logo_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'RIGHT'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ]))
        right_element = logo_table
    else:
        right_html = f"<font size='18'><b><font color='{primary_color_hex}'>НЕВА</font><font color='#D22630'>КОН</font></b></font><br/>{settings.get('address', '')}"
        right_element = Paragraph(right_html, style_right)
    
    header_data = [
        [Paragraph(left_html, style_left_header), right_element]
    ]
    header_table = Table(header_data, colWidths=[250, 265])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 5))
    
    line_table = Table([[""]], colWidths=[515], rowHeights=[2])
    line_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), primary_color),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(line_table)
    story.append(Spacer(1, 8))
    
    contacts_html = f"{settings.get('contacts', '')}<br/><b><font color='{primary_color_hex}'>{settings.get('website', '')}</font></b><br/><i><font color='#718096'>{settings.get('services', '')}</font></i>"
    story.append(Paragraph(contacts_html, style_normal))
    story.append(Spacer(1, 15))
    
    date_str = customer_data.get("date", datetime.now().strftime("%d.%m.%Y"))
    if re.match(r'\d{4}-\d{2}-\d{2}', date_str):
        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        date_str = date_obj.strftime("%d.%m.%Y")
    
    info_data = [
        [Paragraph("<b>Дата:</b>", ParagraphStyle('lbl', parent=style_normal, textColor=primary_color)), Paragraph(date_str, style_normal)],
        [Paragraph("<b>Тел./факс:</b>", ParagraphStyle('lbl', parent=style_normal, textColor=primary_color)), Paragraph(customer_data.get("phone", ""), style_normal)],
        [Paragraph("<b>Кому:</b>", ParagraphStyle('lbl', parent=style_normal, textColor=primary_color)), Paragraph(customer_data.get("to", ""), style_normal)],
        [Paragraph("<b>От:</b>", ParagraphStyle('lbl', parent=style_normal, textColor=primary_color)), Paragraph(customer_data.get("from", ""), style_normal)],
        [Paragraph("<b>Объект:</b>", ParagraphStyle('lbl', parent=style_normal, textColor=primary_color)), Paragraph(customer_data.get("object", ""), style_normal)],
    ]
    
    info_table = Table(info_data, colWidths=[80, 435])
    info_table.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 1, primary_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F7FAFC')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 20))
    story.append(Paragraph("<u>КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ</u>", style_title))
    story.append(Spacer(1, 12))
    
    intro_p = (
        "Благодарим Вас за обращение в нашу фирму по вопросам вентиляции и "
        "кондиционирования воздуха. Для создания комфортных температурных условий "
        "в Вашем помещении предлагаем установить следующий вид кондиционеров:"
    )
    story.append(Paragraph(intro_p, style_normal))
    story.append(Spacer(1, 12))
    
    models = config.load_models()
    
    # Проверяем ВСЕ варианты ДО начала генерации
    missing_models = []
    for var in variants:
        m_key = var["model_key"]
        if m_key not in models:
            missing_models.append(m_key)
    
    if missing_models:
        raise Exception(
            f"Модели не найдены в базе данных: {', '.join(missing_models)}.\n"
            f"Добавьте эти модели во вкладке 'База моделей' перед генерацией документа."
        )
    
    for idx, var in enumerate(variants, 1):
        m_key = var["model_key"]
        model = models[m_key]
        
        story.append(Paragraph(f"<b>{idx} вариант</b>", style_h2))
        
        desc_text = model.get("description", "Сплит-система настенного типа")
        p_desc_style = ParagraphStyle(
            'CyrDescStyle',
            parent=style_normal,
            fontName=font_reg,
            fontSize=10,
            leading=12,
            textColor=colors.HexColor('#4A5568')
        )
        story.append(Paragraph(f"<i>{desc_text}</i>", p_desc_style))
        story.append(Spacer(1, 5))
        
        headers = [
            "Модель", "Уровень шума, дБ", "Произв. по холоду, кВт", "Произв. по теплу, кВт",
            "Потреб. мощн., кВт", "Габариты внутр. блока, мм", "Кол-во, шт.", "Цена, руб.", "Итого, руб."
        ]
        
        qty = var["quantity"]
        price = model["price"]
        total_equip = price * qty
        
        row_vals = [
            m_key,
            model["noise_level"],
            model["cooling_kw"],
            model["heating_kw"],
            model["power_kw"],
            model["dimensions"],
            str(qty),
            f"{price:,}".replace(",", " "),
            f"{total_equip:,}".replace(",", " ")
        ]
        
        table_data = []
        hdr_style = ParagraphStyle('HdrStyle', parent=style_normal, fontSize=8, leading=10, fontName=font_bold_reg, textColor=colors.white)
        header_row = [Paragraph(f"<b>{h}</b>", hdr_style) for h in headers]
        table_data.append(header_row)
        
        val_style = ParagraphStyle('ValStyle', parent=style_normal, fontSize=8, leading=10)
        value_row = [Paragraph(v, val_style) for v in row_vals]
        table_data.append(value_row)
        
        col_widths = [95, 45, 50, 50, 50, 65, 35, 60, 65]
        var_table = Table(table_data, colWidths=col_widths)
        var_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 1, primary_color),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BACKGROUND', (0,0), (-1,0), primary_color),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#F8FAFC')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(var_table)
        story.append(Spacer(1, 10))
        
        total = total_equip + var["installation_price"] + var["approval_price"]
        
        style_total_lbl = ParagraphStyle(
            'TotLbl', parent=style_normal, fontName=font_bold_reg, fontSize=11, textColor=accent_color
        )
        style_total_val = ParagraphStyle(
            'TotVal', parent=style_right, fontName=font_bold_reg, fontSize=11, textColor=accent_color
        )
        
        calc_details = [
            [Paragraph("Монтаж кондиционера", style_normal), Paragraph(f"{var['installation_price']:,} руб.".replace(",", " "), style_right)],
            [Paragraph("Согласование места размещения наружного блока", style_normal), Paragraph(f"{var['approval_price']:,} руб.".replace(",", " "), style_right)],
            [Paragraph("<b>ИТОГО</b>", style_total_lbl), Paragraph(f"<b>{total:,} руб.</b>".replace(",", " "), style_total_val)]
        ]
        
        calc_table = Table(calc_details, colWidths=[355, 160])
        calc_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        story.append(calc_table)
        story.append(Spacer(1, 8))
        
        sep_table = Table([[""]], colWidths=[515], rowHeights=[1])
        sep_table.setStyle(TableStyle([
            ('LINEABOVE', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
        ]))
        story.append(sep_table)
        story.append(Spacer(1, 8))
    
    story.append(Spacer(1, 10))
    story.append(Paragraph("С уважением,", style_normal))
    
    style_sig = ParagraphStyle(
        'CyrSig',
        parent=style_normal,
        fontName=font_bold_reg,
        textColor=primary_color
    )
    story.append(Paragraph(settings.get("manager_signature", "Толстиков Игорь."), style_sig))
    
    try:
        doc_tpl.build(story)
    except PermissionError:
        raise PermissionError(f"Не удалось сохранить PDF-файл '{os.path.basename(filename)}'. Закройте его, если он открыт в просмотрщике PDF, и попробуйте снова.")
    except Exception as e:
        raise Exception(f"Ошибка при создании PDF: {str(e)}")